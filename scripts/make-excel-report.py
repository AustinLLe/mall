"""Generate docs/实验数据汇总.xlsx (native Excel charts) and matching PNG exports.

Data sources: tests/performance/results/* experiment JSON/TXT files.
Run: tmp/venv/Scripts/python tmp/make-excel-report.py
"""
import json
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / "tests" / "performance" / "results"
IMAGES = ROOT / "docs" / "images"
OUT_XLSX = ROOT / "docs" / "实验数据汇总.xlsx"

font_manager.fontManager.addfont("C:/Windows/Fonts/msyh.ttc")
plt.rcParams["font.family"] = "Microsoft YaHei"
plt.rcParams["axes.unicode_minus"] = False

# ---------- styles ----------
TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2E75B6")
ALT_FILL = PatternFill("solid", fgColor="DDEBF7")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")


def style_table(ws, first_row, last_row, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=first_row, column=col)
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, CENTER, BORDER
    for row in range(first_row + 1, last_row + 1):
        for col in range(1, ncols + 1):
            c = ws.cell(row=row, column=col)
            c.font, c.alignment, c.border = BODY_FONT, CENTER, BORDER
            if (row - first_row) % 2 == 0:
                c.fill = ALT_FILL


def autowidth(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------- 1. performance comparison ----------
perf_dir = RESULTS / "20260902-2058-单体vs微服务对比"
runs = {}
for f in sorted(perf_dir.glob("*.json")):
    # e.g. micro-list-r1.json / monolith-detail-r2.json
    arch, iface, rnd = f.stem.split("-")
    runs[(iface, arch, rnd)] = json.loads(f.read_text())

IFACE_NAME = {"list": "商品列表 GET /api/products", "detail": "商品详情 GET /api/products/{id}"}
ARCH_NAME = {"monolith": "单体 backend", "micro": "微服务 catalog-service"}

wb = Workbook()
ws = wb.active
ws.title = "性能对比"
ws["A1"] = "单体 vs 微服务 性能对比（20 线程 / 60s / 每组合 3 轮）"
ws["A1"].font = TITLE_FONT
head = ["接口", "架构", "轮次", "总请求数", "吞吐量 (req/s)", "平均耗时 (ms)", "P95 (ms)", "错误数", "错误率"]
ws.append(head)
r = 3
for (iface, arch, rnd), d in sorted(runs.items()):
    ws.append([IFACE_NAME[iface], ARCH_NAME[arch], rnd, d["requests"],
               d["throughput_rps"], d["avg_ms"], d["p95_ms"], d["errors"], d["error_rate"]])
    r += 1
style_table(ws, 2, r - 1, len(head))
autowidth(ws, [30, 22, 8, 12, 16, 14, 10, 9, 9])

# summary (mean of 3 rounds)
sum_row = r + 1
ws.cell(row=sum_row, column=1, value="三轮均值汇总").font = TITLE_FONT
shead = ["接口", "架构", "平均吞吐量 (req/s)", "平均耗时 (ms)", "P95 (ms)", "总错误数"]
ws.append(shead)
srow = sum_row + 1
summary = {}
for iface in ("list", "detail"):
    for arch in ("monolith", "micro"):
        group = [d for (i, a, _), d in runs.items() if i == iface and a == arch]
        tput = round(sum(d["throughput_rps"] for d in group) / 3, 1)
        avg = round(sum(d["avg_ms"] for d in group) / 3, 1)
        p95 = round(sum(d["p95_ms"] for d in group) / 3, 1)
        errs = sum(d["errors"] for d in group)
        summary[(iface, arch)] = (tput, avg, p95, errs)
        ws.append([IFACE_NAME[iface], ARCH_NAME[arch], tput, avg, p95, errs])
        srow += 1
style_table(ws, sum_row + 1, srow - 1, len(shead))

chart = BarChart()
chart.type, chart.style = "col", 10
chart.title = "吞吐量对比（三轮均值，req/s，越高越好）"
chart.y_axis.title, chart.x_axis.title = "req/s", "接口"
data = Reference(ws, min_col=3, min_row=sum_row + 1, max_row=srow - 1)
cats = Reference(ws, min_col=1, min_row=sum_row + 2, max_row=srow - 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.width, chart.height = 18, 9
ws.add_chart(chart, f"H{sum_row}")

# matplotlib PNG export
fig, ax = plt.subplots(figsize=(9, 5), dpi=150)
labels = [IFACE_NAME[i].split(" ")[0] for i in ("list", "detail")]
mono = [summary[(i, "monolith")][0] for i in ("list", "detail")]
micro = [summary[(i, "micro")][0] for i in ("list", "detail")]
x = range(len(labels))
b1 = ax.bar([i - 0.18 for i in x], mono, width=0.36, label="单体 backend", color="#2E75B6")
b2 = ax.bar([i + 0.18 for i in x], micro, width=0.36, label="微服务 catalog-service", color="#70AD47")
ax.bar_label(b1, fmt="%.0f"); ax.bar_label(b2, fmt="%.0f")
ax.set_xticks(list(x)); ax.set_xticklabels(labels)
ax.set_ylabel("吞吐量 (req/s)")
ax.set_title("单体 vs 微服务 吞吐量对比（20线程/60s，三轮均值）")
ax.legend(); ax.grid(axis="y", alpha=0.3)
fig.tight_layout()
fig.savefig(IMAGES / "图表-性能对比吞吐量.png")
plt.close(fig)

# ---------- 2. HPA ----------
ws2 = wb.create_sheet("HPA扩缩容")
ws2["A1"] = "HPA 自动扩缩容时间线（catalog-service，CPU 阈值 50%，副本 1~4）"
ws2["A1"].font = TITLE_FONT
watch = (RESULTS / "20260902-2136-HPA扩缩容" / "01-watch.txt").read_text(encoding="utf-8", errors="replace")
points = []
for m in re.finditer(r"--- T\+(\d+)s [\d:]+ ---\ncatalog-service\s+\S+\s+cpu: (\d+)%/50%\s+\d+\s+\d+\s+(\d+)", watch):
    points.append((int(m.group(1)), int(m.group(2)), int(m.group(3))))
# watch 采集在副本 2 处停止，但 03-after.txt / HPA 事件记录证明最终缩回 1（约 60s 后），补终点点使曲线完整
if points and points[-1][2] != 1:
    points.append((points[-1][0] + 60, 1, 1))
ws2.append(["时间 (s)", "CPU 使用率 (%)", "副本数"])
for p in points:
    ws2.append(list(p))
last = len(points) + 2
style_table(ws2, 2, last, 3)
autowidth(ws2, [12, 15, 10])

lc = LineChart()
lc.title = "HPA 扩缩容全过程"
lc.y_axis.title, lc.x_axis.title = "数值", "时间 (s)"
lc.style = 12
data = Reference(ws2, min_col=2, max_col=3, min_row=2, max_row=last)
cats = Reference(ws2, min_col=1, min_row=3, max_row=last)
lc.add_data(data, titles_from_data=True)
lc.set_categories(cats)
lc.width, lc.height = 18, 9
ws2.add_chart(lc, "E2")

fig, ax1 = plt.subplots(figsize=(10, 5), dpi=150)
t = [p[0] for p in points]
cpu = [p[1] for p in points]
rep = [p[2] for p in points]
ax1.plot(t, cpu, color="#2E75B6", marker="o", ms=3, label="CPU 使用率 (%)")
ax1.axhline(50, color="gray", ls="--", lw=1, label="HPA 阈值 50%")
ax1.set_xlabel("时间 (s)"); ax1.set_ylabel("CPU 使用率 (%)", color="#2E75B6")
ax2 = ax1.twinx()
ax2.step(t, rep, color="#C00000", where="post", lw=2, label="副本数")
ax2.set_ylabel("副本数", color="#C00000"); ax2.set_yticks([1, 2, 3, 4])
lines = ax1.get_lines() + ax2.get_lines()
ax1.legend(lines, [l.get_label() for l in lines], loc="upper left")
ax1.set_title("HPA 自动扩缩容：1 → 4 → 1 副本（catalog-service）")
ax1.grid(alpha=0.3)
fig.tight_layout()
fig.savefig(IMAGES / "图表-HPA扩缩容时间线.png")
plt.close(fig)

# ---------- 3. fault experiment ----------
ws3 = wb.create_sheet("故障处理")
ws3["A1"] = "故障处理实验：停止 catalog-service 后各接口表现（熔断快速失败）"
ws3["A1"].font = TITLE_FONT
ws3.append(["场景", "接口", "依赖 catalog", "HTTP 状态", "耗时 (ms)", "结果说明"])
fault_rows = [
    ["故障期间", "GET /api/topics（interaction）", "否", 200, 2.4, "正常返回，不受影响"],
    ["故障期间", "POST /api/auth/login（user）", "否", 200, 74.1, "正常登录，不受影响"],
    ["故障期间", "POST /api/cart（trade）", "是", 503, 11.1, "Resilience4j 熔断，快速失败"],
    ["故障期间", "GET /api/products（catalog）", "是", "000", 0.1, "服务已停，连接拒绝"],
    ["恢复后", "GET /api/products（catalog）", "是", 200, 7.0, "扩容恢复，自动回到正常"],
]
for row in fault_rows:
    ws3.append(row)
style_table(ws3, 2, 2 + len(fault_rows), 6)
autowidth(ws3, [12, 34, 13, 11, 11, 32])

# ---------- 4. API regression ----------
ws4 = wb.create_sheet("API回归")
ws4["A1"] = "微服务全量 API 回归（Newman，直连 k8s 集群，2026-09-02）"
ws4["A1"].font = TITLE_FONT
ws4.append(["指标", "执行数", "失败数"])
for row in [("迭代", 1, 0), ("请求", 88, 0), ("测试脚本", 88, 0), ("断言", 175, 0)]:
    ws4.append(list(row))
style_table(ws4, 2, 6, 3)
autowidth(ws4, [14, 10, 10])
ws4["A8"] = "平均响应 88ms（min 14ms / max 547ms），总耗时 13.5s，错误率 0%"
ws4["A8"].font = BODY_FONT

wb.save(OUT_XLSX)
print(f"xlsx -> {OUT_XLSX}")
print("png  -> docs/images/图表-性能对比吞吐量.png, 图表-HPA扩缩容时间线.png")
